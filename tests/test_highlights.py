"""Checks against the highlight annotations drawn on the two sample documents.

Every expected value here was read off the highlighted PDF by hand.
"""

import sys
from pathlib import Path

import pdfplumber
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from boe_extraction.extract import extract_with_highlights
from boe_extraction.highlights import read as read_highlights

SAMPLES = Path(__file__).resolve().parents[1] / "samples"
COURIER = SAMPLES / "FBA15M6L9KGF01_courier_cbe_xiv.pdf"
STANDARD = SAMPLES / "BOE_3141398_icegate.pdf"
XIII = SAMPLES / "FBA15M1ZPS2Y01_cbe_xiii_highlighted.pdf"


@pytest.fixture(scope="module")
def courier():
    return extract_with_highlights(COURIER)


@pytest.fixture(scope="module")
def standard():
    return extract_with_highlights(STANDARD)


def _values(fields):
    return {f.label: f.value for f in fields}


def test_every_highlight_is_read(courier, standard):
    assert len(courier[2]) == 55
    assert len(standard[2]) == 84


def test_highlight_text_is_not_merged_across_rows():
    """A crop would return "Port Code" over "INNSA1" as "IPNoNrtS CAo1d"."""
    with pdfplumber.open(STANDARD) as pdf:
        texts = [h.text for h in read_highlights(pdf) if h.page == 0]
    assert "Port Code" in texts
    assert "INNSA1" in texts


def test_courier_highlighted_fields(courier):
    values = _values(courier[1])
    assert values["CBEXIV Number"] == "CBEXIV_DEL_2026-2027_2808_10570"
    assert values["Import export Code"] == "AAFCV5265N"
    assert values["KYC ID"] == "29AAFCV5265N1ZK"
    assert values["Import General Manifest (IGM) Number"] == "3129351"
    assert values["Master Airway Bill (MAWB) Number"] == "31296805785"
    assert values["House Airway Bill (HAWB) Number"] == "FBA15M6L9KGF01"
    assert values["Invoice Number"] == "FBA15M6L9KGF01"
    assert values["Invoice Value"] == "1080"
    assert values["Terms of Invoice"] == "CIF"
    # Per-item fields live on the Line Items sheet, not among the document's.
    assert "CTSH" not in values
    assert "Assessable Value" not in values


def test_section_banners_are_not_reported_as_fields(courier):
    """A highlight over a heading or a multi-row table names no field.

    Its figures are reported against each line item instead, so it is left out
    of the document's fields and kept only in Highlights (raw).
    """
    assert not [f for f in courier[1] if f.label == "(as highlighted)"]
    raw = " ".join(h.text for h in courier[2])
    assert "1 BCD 20 0 0 3908" in raw


def test_single_row_tables_become_fields(courier):
    """A bold heading row over one row of figures is a record, so read it."""
    values = _values(courier[1])
    assert values["Current Status of the CBE"] == "OOC ISSUED on 29-08-2026 07:11"
    assert values["TR-6 Challan Number"] == "2908387251"
    assert values["Total Amount"] == "45833"
    assert values["Challan Date"] == "29/08/2026"
    assert values["Airlines"] == "Indigo Airlines"
    assert values["Flight No."] == "6E 1074"
    assert values["Airport Of Arrival"] == "DEL"
    assert values["Date Of Arrival"] == "27/08/2026"


def test_sections_tell_repeated_labels_apart(courier):
    """The importer, the supplier and the broker each have a Name."""
    names = {f.key: f.value for f in courier[1] if f.label == "Name"}
    assert names["PARTICULARS OF THE IMPORTER · Name"] == "VALUECART PRIVATE LIMITED"
    assert names["SUPPLIER DETAILS · Name"] == "GATI HONG KONG LIMITED"
    assert names["BROKER/ AGENT DETAILS · Name"] == "KBR INTERNATIONAL LOGISTICS"


def test_standard_highlighted_summary_totals(standard):
    """The Part I summary is a heading row over a value row."""
    values = _values(standard[1])
    assert values["1.BCD"] == "84001.2"
    assert values["3.SWS"] == "8400.3"
    assert values["7.IGST"] == "117434"
    assert values["18.TOT.ASS VAL"] == "560008"
    assert values["14.TOTAL DUTY"] == "209835"
    assert values["19.TOT. AMOUNT"] == "209835"


def test_standard_highlighted_header_fields(standard):
    values = _values(standard[1])
    assert values["Port Code"] == "INNSA1"
    assert values["BE No"] == "3141398"
    assert values["BE Date"] == "14/08/2026"
    assert values["G.WT (KGS)"] == "691"      # the barcode beside it is dropped
    assert values["GSTIN"] == "27AAFCV5265N1ZO"
    assert values["TYPE"] == "G"
    assert values["IEC"] == "AAFCV5265N"
    assert values["Br"] == "1"
    assert values["CB CODE"] == "AAHFS9149KCH001"
    assert values["AD CODE"] == "6480001"
    assert values["13.COUNTRY OF ORIGIN"] == "CHINA"
    assert values["15.PORT OF LOADING"] == "Shekou"
    assert values["2.CB NAME"] == "INTERLINK SHIPPING & CLEARING"

    # Part II repeats per invoice, so its figures are reported per invoice on
    # the Invoices and Line Items sheets, not once at document level.
    for label in ("2.INVOICE NO. & DT.", "3.PURCHASE ORDER NO & DT",
                  "1.INV VALUE", "2.FREIGHT", "14.Cur", "15.Term",
                  "14.ASS. VALUE"):
        assert label not in values
    invoice = standard[0].invoices[0]
    assert (invoice.number, invoice.invoice_value, invoice.currency) == (
        "FBA15M13GSD3", 3050.3, "USD")
    assert round(sum(i.assessable_value for i in invoice.items), 2) == 292981.32


def test_a_name_and_address_is_read_in_full(standard):
    """The heading sits indented over a block that starts further left.

    Reading only the line beneath it returned the name without the address.
    """
    values = _values(standard[1])
    assert values["1.IMPORTER NAME & ADDRESS"] == (
        "VALUECART PRIVATE LIMITED, FLAT No-4, G. S. TOWERS, OPPOSITE H, "
        "BIBWEWADI, PUNE, PUNE, 411037")
    assert values["1.BUYER'S NAME & ADDRESS"].endswith("411037")
    assert values["3.SUPPLIER NAME & ADDRESS"].startswith(
        "GATI HONG KONG LIMITED, FLAT/RM C1303")
    # This document names no third party; the block under that heading is
    # empty, and must not borrow the supplier's.
    assert values["4.THIRD PARTY NAME & ADDRESS"] == ""


def test_every_item_is_reported_on_line_items(tmp_path, courier):
    """A field highlighted on item 1 is wanted for all of them."""
    from openpyxl import load_workbook

    from boe_extraction.excel_writer import write_highlighted

    boe, fields, highlights = courier
    path = write_highlighted(boe, fields, highlights, tmp_path / "h.xlsx")
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Highlighted Fields", "Line Items",
                                   "Highlights (raw)"]
    sheet = workbook["Line Items"]
    header = [c.value for c in sheet[1]]
    assert sheet.max_row == 5  # header plus four items

    column = header.index("Assessable Value") + 1
    assert [sheet.cell(row=r, column=column).value for r in range(2, 6)] == [
        19541.25, 26055, 6513.75, 52110]


def test_workbook_records_every_highlight_verbatim(tmp_path, standard):
    from openpyxl import load_workbook

    from boe_extraction.excel_writer import write_highlighted

    boe, fields, highlights = standard
    path = write_highlighted(boe, fields, highlights, tmp_path / "h.xlsx")
    sheet = load_workbook(path)["Highlights (raw)"]
    assert sheet.max_row == len(highlights) + 1


def test_each_document_gets_its_own_workbook(tmp_path):
    """No sheet ever mixes two documents: one workbook per bill of entry."""
    from openpyxl import load_workbook

    from boe_extraction.excel_writer import write_mandatory
    from boe_extraction.extract import extract_document

    # Only the courier form highlights a field that varies per line item, so
    # only it earns a Mandatory Item Fields sheet.
    expected = {COURIER: ["Mandatory Fields", "Mandatory Item Fields",
                          "Line Items", "Field Checklist", "Highlights (raw)"],
                STANDARD: ["Mandatory Fields", "Line Items", "Field Checklist",
                           "Highlights (raw)"]}
    for pdf in (COURIER, STANDARD):
        document = extract_document(pdf)
        path = write_mandatory(document, tmp_path / f"{pdf.stem}.xlsx")
        workbook = load_workbook(path)
        assert workbook.sheetnames == expected[pdf]
        # A Document column is what a combined workbook needs; these have none.
        for name in workbook.sheetnames:
            assert workbook[name]["A1"].value != "Document"
        assert workbook["Mandatory Fields"].max_row == 2      # header plus one


def test_mandatory_workbook_uses_the_fields_as_columns(tmp_path):
    """The highlights are the required-field list, so they are the columns."""
    from openpyxl import load_workbook

    from boe_extraction.excel_writer import write_mandatory
    from boe_extraction.extract import extract_document

    def header_and_row(pdf):
        path = write_mandatory(extract_document(pdf), tmp_path / f"{pdf.stem}.xlsx")
        sheet = load_workbook(path)["Mandatory Fields"]
        header = [c.value for c in sheet[1]]
        row = next(sheet.iter_rows(min_row=2, values_only=True))
        return header, row

    header, row = header_and_row(COURIER)
    assert header[:2] == ["Form Type", "BE Number"]
    assert row[header.index("CBEXIV Number")] == "CBEXIV_DEL_2026-2027_2808_10570"
    assert row[header.index("TR-6 Challan Number")] == "2908387251"

    header, row = header_and_row(STANDARD)
    assert row[header.index("1.BCD")] == "84001.2"
    assert row[header.index("EXCHANGE RATE")] == "1 USD=96.05INR"
    assert "PARTICULARS OF THE IMPORTER · Name" in header_and_row(COURIER)[0]


def test_item_rows_cover_every_item_of_the_document(tmp_path):
    from openpyxl import load_workbook

    from boe_extraction.excel_writer import write_mandatory
    from boe_extraction.extract import extract_document

    for pdf, count in ((COURIER, 4), (STANDARD, 9)):
        path = write_mandatory(extract_document(pdf), tmp_path / f"{pdf.stem}.xlsx")
        sheet = load_workbook(path)["Line Items"]
        assert sheet.max_row == count + 1


def test_an_invoice_workbook_holds_only_that_invoice(tmp_path):
    """BOE 3141398 carries two invoices; neither may reach the other's file."""
    from openpyxl import load_workbook

    from boe_extraction.excel_writer import write_mandatory
    from boe_extraction.extract import extract_document

    document = extract_document(STANDARD)
    assert [i.number for i in document.boe.invoices] == ["FBA15M13GSD3",
                                                         "FBA15M16XHDH"]
    for invoice, count in zip(document.boe.invoices, (5, 4)):
        path = write_mandatory(document, tmp_path / f"{invoice.number}.xlsx",
                               invoice)
        workbook = load_workbook(path)
        assert workbook["Line Items"].max_row == count + 1
        # One invoice per file, so no sheet needs to name which invoice.
        assert workbook["Line Items"]["A1"].value != "Invoice"


def test_the_item_table_is_reported_on_line_items(standard):
    """Part II's item columns are per item, so they belong on Line Items."""
    values = _values(standard[1])
    for label in ("1.S NO.", "2.CTH", "3.DESCRIPTION", "4.UNIT PRICE",
                  "5.QUANTITY", "6.UQC", "7.AMOUNT"):
        assert label not in values
    first = standard[0].all_items()[0]
    assert first.hs_code == "39269099"
    assert first.description == "X002HTKLZJ MACBOOK PRO 16 INCH CASE - PC"
    assert (first.unit_price, first.quantity, first.unit_of_measure) == (4.62, 40, "NOS")


def test_a_blank_field_does_not_borrow_an_item_value(courier):
    """The courier form carries a Country of Origin at both levels.

    Only the document's own is reported here -- the item's is on Line Items --
    and a blank one must stay blank rather than take the item's value.
    """
    origins = [f for f in courier[1] if f.label == "Country of Origin"]
    assert [f.value for f in origins] == [""]
    assert origins[0].section == "SPECIAL REQUESTS"
    assert courier[0].all_items()[0].details["Country of Origin"] == "CHINA"


CBE_XIII_HIGHLIGHTED = SAMPLES / "FBA15M1ZPS2Y01_cbe_xiii_highlighted.pdf"


def test_highlights_resolve_on_a_cbe_xiii():
    """The highlight path had only ever run on the other two forms."""
    boe, fields, highlights = extract_with_highlights(CBE_XIII_HIGHLIGHTED)
    assert boe.form_type == "CBE-XIII"
    assert len(highlights) == 9
    values = _values(fields)
    assert values["Import Export Code"] == "AAFCV5265N"
    assert values["KYC ID"] == "29AAFCV5265N1ZK"
    # Its per-item fields are reported on Line Items instead.
    assert "CTSH" not in values
    assert boe.all_items()[0].hs_code == "42023290"


def test_a_document_total_is_not_replaced_by_an_items_share():
    """A courier form carries an Assessable Value at both levels.

    The consignment's total must survive; only a field inside an item's own
    block may be repaired from that item.
    """
    _, fields, _ = extract_with_highlights(CBE_XIII_HIGHLIGHTED)
    values = _values(fields)
    assert values["Assessable Value"] == "82062.07"   # not item 1's 2844.07
    assert values["Duty(Rs.)"] == "29648"             # not item 1's 1066
    boe, _, _ = extract_with_highlights(CBE_XIII_HIGHLIGHTED)
    assert boe.be_number == "CBEXIII_DEL_2026-2027_2707_14754"


def test_a_value_wrapped_onto_a_second_line_is_joined(standard):
    """The MAWB number is printed as "ONEYSZPGM" over "8034800"."""
    values = _values(standard[1])
    assert values["6.MAWB NO"] == "ONEYSZPGM8034800"
    assert values["8.HAWB NO"] == "GATI26SE0725"
    # A wrapped phrase keeps its space rather than joining up closed.
    assert values["2.CB NAME"] == "INTERLINK SHIPPING & CLEARING"


def test_a_wrapped_heading_is_not_read_as_values(standard):
    """"7.ADV BE 11.FIRST 12. PROV/" wraps onto "(Y/N/P) CHECK FINAL".

    Both lines are heading; the values are further down. Reading the wrapped
    line as values turned the flag "N" into "CHECKN". Those processing flags
    are not reported, but the fields around them must survive intact.
    """
    values = _values(standard[1])
    assert values["1.BE STATUS"] == "FIRST COPY"
    assert values["2.MODE"] == "Sea"
    # A heading whose value sits beside it still resolves.
    assert values["15.PORT OF LOADING"] == "Shekou"
    assert values["16.PORT OF SHIPMENT"] == "Shekou"


def test_one_workbook_gives_each_document_its_own_sheets(tmp_path):
    """A run may share a file, but never a sheet."""
    from openpyxl import load_workbook

    from boe_extraction.excel_writer import write_workbook
    from boe_extraction.extract import extract_document

    documents = [extract_document(COURIER), extract_document(STANDARD)]
    path = write_workbook(documents, tmp_path / "by_document.xlsx")
    workbook = load_workbook(path)
    assert workbook.sheetnames == [
        "Courier CBE-XIV Invoices", "Courier CBE-XIV Fields",
        "Courier CBE-XIV Line Items", "Courier CBE-XIV Highlights",
        "Cargo BOE Invoices", "Cargo BOE Fields", "Cargo BOE Line Items",
        "Cargo BOE Highlights"]

    # No sheet names a document, because no sheet holds more than one.
    for name in workbook.sheetnames:
        assert workbook[name]["A1"].value != "Document"
    assert workbook["Courier CBE-XIV Line Items"].max_row == 5      # 4 items

    # Every invoice row names the bill of entry it was raised under.
    invoices = workbook["Cargo BOE Invoices"]
    assert [c.value for c in invoices[1]][:2] == ["BE No", "Invoice Number"]
    assert [r[:2] for r in invoices.iter_rows(min_row=2, values_only=True)] == [
        ("3141398", "FBA15M13GSD3"), ("3141398", "FBA15M16XHDH")]
    assert workbook["Courier CBE-XIV Invoices"]["A2"].value == (
        "CBEXIV_DEL_2026-2027_2808_10570")
    assert workbook["Cargo BOE Highlights"].max_row == 85
    assert workbook["Courier CBE-XIV Highlights"].max_row == 56

    # The cargo bill's two invoices stay on one sheet, named by a column.
    items = workbook["Cargo BOE Line Items"]
    assert items.max_row == 10                                      # 9 items
    assert items["A1"].value == "Invoice"
    assert {r[0] for r in items.iter_rows(min_row=2, values_only=True)} == {
        "FBA15M13GSD3", "FBA15M16XHDH"}


def test_per_invoice_sheets_split_a_multi_invoice_document(tmp_path):
    from openpyxl import load_workbook

    from boe_extraction.excel_writer import write_workbook
    from boe_extraction.extract import extract_document

    path = write_workbook([extract_document(STANDARD)],
                          tmp_path / "by_invoice.xlsx", per_invoice=True)
    workbook = load_workbook(path)
    # Only sheets that really do cover one invoice are named for it.
    assert workbook.sheetnames == [
        "FBA15M13GSD3 Invoices", "FBA15M13GSD3 Fields",
        "FBA15M13GSD3 Line Items", "FBA15M13GSD3 Highlights",
        "FBA15M16XHDH Invoices", "FBA15M16XHDH Fields",
        "FBA15M16XHDH Line Items", "FBA15M16XHDH Highlights"]
    for name, number in (("FBA15M13GSD3", "FBA15M13GSD3"),
                         ("FBA15M16XHDH", "FBA15M16XHDH")):
        sheet = workbook[f"{name} Invoices"]
        assert sheet.max_row == 2                       # header plus its own
        assert [c.value for c in sheet[1]][:2] == ["BE No", "Invoice Number"]
        assert (sheet["A2"].value, sheet["B2"].value) == ("3141398", number)
    assert workbook["FBA15M13GSD3 Line Items"].max_row == 6        # 5 items
    assert workbook["FBA15M16XHDH Line Items"].max_row == 5        # 4 items
    # Each invoice's sheet stands alone, so it need not name the invoice.
    for name in ("FBA15M13GSD3 Line Items", "FBA15M16XHDH Line Items"):
        assert workbook[name]["A1"].value == "Item Number"


def test_a_sheet_name_stays_within_the_excel_limit(tmp_path):
    """The courier BE numbers are longer than Excel allows in a sheet name."""
    from openpyxl import load_workbook

    from boe_extraction.excel_writer import SHEET_NAME_LIMIT, write_workbook
    from boe_extraction.extract import extract_document

    documents = [extract_document(p) for p in (COURIER, XIII, STANDARD)]
    path = write_workbook(documents, tmp_path / "all.xlsx")
    names = load_workbook(path).sheetnames
    assert len(names) == len(set(names))
    for name in names:
        assert len(name) <= SHEET_NAME_LIMIT


def test_the_schema_is_free_of_duplicates():
    """No field is named twice, and no value appears under two names."""
    from boe_extraction.schema import (CBE_XIII_FIELDS, CBE_XIV_FIELDS,
                                       ITEM_FIELDS)

    for fields in (CBE_XIV_FIELDS, CBE_XIII_FIELDS):
        assert len(fields) == len(set(fields))
    assert len(ITEM_FIELDS) == len(set(ITEM_FIELDS))
    # The columns that repeated another column's value on every item.
    for dropped in ("CTSH", "CETSH", "Description of Goods", "Duty(Rs.)",
                    "Rate of Exchange", "Invoice Number", "Currency of Invoice",
                    "Charge Type", "Charge Amount(in rs.)"):
        assert dropped not in ITEM_FIELDS
    assert ("PARTICULARS OF THE IMPORTER", "BOE Number") not in CBE_XIV_FIELDS


@pytest.mark.parametrize("pdf,form,fields,items",
                         [(COURIER, "CBE-XIV", 55, 4), (XIII, "CBE-XIII", 35, 44),
                          (STANDARD, "ICEGATE BOE", 51, 9)])
def test_the_extract_carries_the_schema_and_nothing_else(tmp_path, pdf, form,
                                                         fields, items):
    from openpyxl import load_workbook

    from boe_extraction.excel_writer import write_schema
    from boe_extraction.extract import schema_extract
    from boe_extraction.schema import DOCUMENT_FIELDS, ITEM_FIELDS

    boe, document, rows = schema_extract(pdf)
    assert boe.form_type == form
    assert [(s, l) for s, l, _ in document] == DOCUMENT_FIELDS[form]
    assert len(document) == fields
    assert len(rows) == items

    path = write_schema(boe, document, rows, tmp_path / "extract.xlsx")
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Document Fields", "Line Items"]
    assert workbook["Document Fields"].max_row == fields + 1
    assert [c.value for c in workbook["Line Items"][1]] == list(ITEM_FIELDS)
    assert workbook["Line Items"].max_row == items + 1


def test_a_wrapped_label_keeps_the_whole_value(tmp_path):
    """The form wraps a label with its colon on the last line.

    "Address of Authorized" over "Courier :" was read as a label "Courier"
    starting at the second line, which lost the first line of the address.
    """
    from boe_extraction.extract import schema_extract

    _, document, _ = schema_extract(XIII)
    values = {label: value for _, label, value in document}
    assert values["Address of Authorized Courier"] == (
        "E 149 GROUND FLOOR WEST PATEL NAGARN/ ANEW DELHIDELHI110008")
    assert values["Name of the Authorized Courier"] == "KBR INTERNATIONAL LOGISTICS"
    # The form prints this label without a colon, so only the schema naming it
    # brings it through.
    assert values["CBE-XIII Number"] == "CBEXIII_DEL_2026-2027_2707_14754"
    # A centred heading with a lower-case qualifier -- "DETAILS OF CRN (if
    # present)" -- used to be read as this empty field's value.
    assert values["Import Using e-Commerce"] == ""


def test_every_schema_item_column_is_filled_where_the_form_fills_it():
    """A column the form answers must not come back empty for every item."""
    from boe_extraction.extract import schema_extract
    from boe_extraction.schema import ITEM_FIELDS

    _, _, rows = schema_extract(XIII)
    filled = {title for index, title in enumerate(ITEM_FIELDS)
              if any(row[index] not in (None, "") for row in rows)}
    for title in ("Invoice", "HS Code", "Description", "Quantity", "Unit Price",
                  "Assessable Value", "BCD Rate", "BCD Amount", "SWS Amount",
                  "IGST Amount", "Duty Amount", "Exchange Rate",
                  "Country of Origin", "Name of Manufacturer", "Invoice Term",
                  "Notification number", "serial number of notification"):
        assert title in filled, title


def test_one_workbook_holds_every_document_on_its_own_sheets(tmp_path):
    """Combined into one file, never into one sheet."""
    from openpyxl import load_workbook

    from boe_extraction.excel_writer import write_schema_combined
    from boe_extraction.extract import schema_extract
    from boe_extraction.schema import ITEM_FIELDS

    extracts = [schema_extract(p) for p in (STANDARD, COURIER, XIII)]
    path = write_schema_combined(extracts, tmp_path / "all.xlsx")
    workbook = load_workbook(path)
    assert workbook.sheetnames == [
        "Cargo BOE Fields", "Cargo BOE Line Items",
        "Courier CBE-XIV Fields", "Courier CBE-XIV Line Items",
        "Courier CBE-XIII Fields", "Courier CBE-XIII Line Items"]
    for name in workbook.sheetnames:
        assert len(name) <= 31
        # No sheet names a document, because no sheet holds more than one.
        assert workbook[name]["A1"].value in ("Section", ITEM_FIELDS[0])
    for name, rows in (("Cargo BOE Fields", 51), ("Cargo BOE Line Items", 9),
                       ("Courier CBE-XIV Fields", 55),
                       ("Courier CBE-XIII Line Items", 44)):
        assert workbook[name].max_row == rows + 1


def test_the_cargo_bill_reads_the_same_with_or_without_highlights():
    """The schema path must not need a reviewer to have marked the PDF."""
    from boe_extraction.extract import extract_with_highlights, schema_extract

    _, highlighted, _ = extract_with_highlights(STANDARD)
    _, fields, _ = schema_extract(STANDARD)
    assert [(f.section, f.label, f.value) for f in highlighted] == fields


def test_duty_amounts_add_up_on_every_item():
    """Total duty is the sum of the heads, which is the form's own check."""
    from boe_extraction.extract import schema_extract
    from boe_extraction.schema import ITEM_FIELDS

    for pdf in (COURIER, XIII, STANDARD):
        _, _, rows = schema_extract(pdf)
        at = {title: index for index, title in enumerate(ITEM_FIELDS)}
        for row in rows:
            heads = sum(row[at[f"{head} Amount"]] or 0
                        for head in ("BCD", "SWS", "IGST", "AIDC", "ADD",
                                     "CHCESS", "CESS", "CMPNSTRY"))
            assert abs(heads - (row[at["Duty Amount"]] or 0)) <= 1.0


def test_a_column_is_read_under_either_form_s_wording():
    """CBE-XIV says "Currency for Unit Price", CBE-XIII "Currency of ...".

    One field, so one column: the other form's wording is an alias, not a
    column of its own.
    """
    from boe_extraction.extract import schema_extract
    from boe_extraction.schema import ITEM_FIELDS

    at = ITEM_FIELDS.index("Currency of Unit Price")
    for pdf in (COURIER, XIII):
        _, _, rows = schema_extract(pdf)
        assert {row[at] for row in rows} == {"USD"}


def test_a_blank_item_column_is_blank_on_the_form():
    """Every column left empty must be one the form does not fill per item.

    The counterpart is either absent from that form, blank on it, or printed
    once at document level -- where it is extracted on the Fields sheet, and
    must not be repeated onto every item row.
    """
    from boe_extraction.extract import schema_extract
    from boe_extraction.schema import ITEM_FIELDS

    # Neither courier form carries these duty heads, nor a manufacturer's
    # address; both leave the licence pair blank.
    absent = {"ADD Rate", "ADD Amount", "CHCESS rate", "CHCESS Amount",
              "CESS rate", "CESS Amount", "License Type", "License Number",
              "Address of Manufacturer", "Discount Amount",
              "Currency of Discount"}
    # CBE-XIV prints these once per invoice rather than once per item.
    at_document_level = {"Number of Packages", "Marks on Packages",
                         "Invoice Value", "Invoice Term", "Landing Charges",
                         "Insurance", "Freight"}

    _, fields, rows = schema_extract(COURIER)
    empty = {title for index, title in enumerate(ITEM_FIELDS)
             if all(row[index] in (None, "") for row in rows)}
    assert empty == absent | at_document_level

    # ... and those really are on the CBE-XIV's own Fields sheet.
    values = {label: value for _, label, value in fields}
    assert values["Number of Packages"] == "1"
    assert values["Invoice Value"] == "1080"
    assert values["Terms of Invoice"] == "CIF"

    _, _, rows = schema_extract(XIII)
    empty = {title for index, title in enumerate(ITEM_FIELDS)
             if all(row[index] in (None, "") for row in rows)}
    assert empty == absent          # CBE-XIII prints the rest per item


def test_the_run_produces_one_workbook_per_family(tmp_path):
    """The cargo bills in one file, the courier bills in another."""
    from openpyxl import load_workbook

    from boe_extraction.cli import main

    assert main([str(STANDARD), str(COURIER), str(XIII), "--schema",
                 "-o", str(tmp_path)]) == 0
    assert sorted(p.name for p in tmp_path.glob("*.xlsx")) == [
        "Cargo_BOE_extract.xlsx", "Courier_BOE_extract.xlsx"]

    cargo = load_workbook(tmp_path / "Cargo_BOE_extract.xlsx")
    assert cargo.sheetnames == ["Cargo BOE Fields", "Cargo BOE Line Items"]
    assert cargo["Cargo BOE Fields"].max_row == 52          # 51 fields
    assert cargo["Cargo BOE Line Items"].max_row == 10      # 9 items

    courier = load_workbook(tmp_path / "Courier_BOE_extract.xlsx")
    assert courier.sheetnames == [
        "Courier CBE-XIV Fields", "Courier CBE-XIV Line Items",
        "Courier CBE-XIII Fields", "Courier CBE-XIII Line Items"]
    assert courier["Courier CBE-XIV Line Items"].max_row == 5
    assert courier["Courier CBE-XIII Line Items"].max_row == 45


def test_a_family_with_no_documents_writes_no_workbook(tmp_path):
    from boe_extraction.cli import main

    assert main([str(COURIER), "--schema", "-o", str(tmp_path)]) == 0
    assert [p.name for p in tmp_path.glob("*.xlsx")] == ["Courier_BOE_extract.xlsx"]


SECOND_CARGO = SAMPLES / "BOE_3036066_icegate.pdf"


def test_a_second_cargo_bill_reads_its_own_figures():
    """A bill the schema was not written against still reconciles."""
    from boe_extraction.extract import schema_extract, verify_document
    from boe_extraction.verify import report

    boe, fields, items = schema_extract(SECOND_CARGO)
    assert (boe.form_type, boe.be_number) == ("ICEGATE BOE", "3036066")
    assert len(fields) == 51 and len(items) == 12

    _, checks = verify_document(SECOND_CARGO)
    _, passed = report(checks, boe.form_type)
    assert passed


def test_part_two_is_the_same_section_however_many_invoices():
    """Part II names the invoices it covers, and that is not its name.

    The heading reads "(Invoice 1 2 )" on a bill with two invoices and
    "(Invoice 1 1 )" on a bill with one, which put the section out of reach of
    a schema written against the other.
    """
    from boe_extraction.extract import schema_extract

    for pdf in (STANDARD, SECOND_CARGO):
        _, fields, _ = schema_extract(pdf)
        sections = {section for section, _, _ in fields}
        assert "PART - II - INVOICE & VALUATION DETAILS" in sections
        assert not [s for s in sections if "(Invoice" in s]


def test_two_address_blocks_side_by_side_stay_apart():
    """The headings are indented over blocks that start further left.

    Cutting the blocks where the headings sit gave the left block the first
    words of the right one, and left the right block missing them.
    """
    from boe_extraction.extract import schema_extract

    _, fields, _ = schema_extract(SECOND_CARGO)
    values = {label: value for _, label, value in fields}
    assert values["1.BUYER'S NAME & ADDRESS"] == (
        "VALUECART PRIVATE LIMITED, NO 2ND FLOOR , 1/1 VINAYAKA TOWERS, "
        "1ST CROSS , GANDHINAGAR, Bangalore, BANGALORE, 560009")
    assert values["2.SELLER'S NAME & ADDRESS"] == (
        "CREATIVE TOOLS HK COMPANY LIMITED, UNIT 03 3/F SEAPOWERCENTRE "
        "NO.73-77, LEI MUK ROAD KWAI CHUNG NT, HK")
    assert values["3.SUPPLIER NAME & ADDRESS"].startswith(
        "VALOCITYCRAFT (THAILAND) CO.,LTD, NO. 369/37 MOO6")
    assert values["4.THIRD PARTY NAME & ADDRESS"].endswith("Hong Kong")

    # The block on page 1 shares its lines with the broker's name and the AD
    # code, and must take neither.
    assert values["1.IMPORTER NAME & ADDRESS"].endswith("BANGALORE, 560009")
    assert values["2.CB NAME"] == "CLASSIC CLEARING & FORWARDING PVT.LTD."
